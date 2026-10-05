# -*- coding: utf-8 -*-
"""Importa CSV/Excel para o banco. Uso: py importar_planilha.py ARQUIVO [--dry-run]"""
import asyncio
import csv
import sys
from pathlib import Path

try:
    from dotenv import load_dotenv
    load_dotenv()
except Exception:
    pass

from db import fetch_one, execute, init_db, USE_POSTGRES, DB_PATH, DATABASE_URL
# melhoria v20: validacao unificada (antes duplicava as regras do main.py)
from validacao import validar_codigo, validar_email, normalizar_celular, normalizar_status

ALIASES = {
    "destinatario": {"nome", "destinatario", "cliente"},
    "codigo": {"codigo", "cod", "rastreio", "objeto"},
    "celular": {"celular", "whatsapp", "whats", "fone", "telefone", "tel", "cel"},
    "email": {"email", "mail"},
    "descricao": {"descricao", "observacoes", "obs", "desc"},
    "status": {"status", "situacao"},
}


def norm_key(s):
    s = str(s or "").strip().lower().replace("_", " ")
    return " ".join("".join(c for c in s if c.isalnum() or c == " ").split())


def map_header(header):
    idx = {}
    for i, h in enumerate(header):
        k = norm_key(h).replace(" ", "")
        for campo, nomes in ALIASES.items():
            flat = {n.replace(" ", "") for n in nomes}
            if k in flat:
                idx.setdefault(campo, i)
    return idx if "codigo" in idx else None


def validar_codigo(codigo):
    from validacao import validar_codigo as _v
    return _v(codigo)


def validar_email(email):
    from validacao import validar_email as _v
    return _v(email)


def normalizar_celular(valor):
    from validacao import normalizar_celular as _n
    return _n(valor)


def ler_csv(path):
    raw = path.read_bytes()
    t = raw.decode("utf-8-sig") if raw.startswith(b"\xef\xbb\xbf") else raw.decode("utf-8", errors="replace")
    prim = next((l for l in t.splitlines() if l.strip()), "")
    sep = ";" if prim.count(";") > prim.count(",") and prim.count(";") >= prim.count("\t") else ("\t" if prim.count("\t") > prim.count(",") else ",")
    return [l for l in csv.reader(t.splitlines(), delimiter=sep) if any(str(c or "").strip() for c in l)]


def ler_excel(path):
    try:
        import openpyxl
    except ImportError:
        print("Falta 'openpyxl'. Instale: py -m pip install openpyxl")
        sys.exit(1)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sh = wb[wb.sheetnames[0]]
    out = []
    for row in sh.iter_rows(values_only=True):
        vals = [("" if c is None else str(c).strip()) for c in row]
        if any(vals):
            out.append(vals)
    return out


def linhas_para_itens(matriz):
    if not matriz:
        return [], ["arquivo vazio"]
    mapa = map_header(matriz[0])
    inicio, avisos = 1, []
    if not mapa:
        mapa = {"destinatario": 0, "codigo": 1, "celular": 2, "email": 3, "descricao": 4}
        if len(matriz[0]) >= 6:
            mapa["status"] = 5
        avisos.append("sem cabecalho — usando ordem Nome,Codigo,Celular,Email,Descricao[,Status]")
        inicio = 0
    itens = []
    for n, lin in enumerate(matriz[inicio:], inicio + 1):
        def get(c):
            i = mapa.get(c)
            return str(lin[i] if i is not None and i < len(lin) else "").strip()
        itens.append({"destinatario": get("destinatario"), "codigo": get("codigo"),
                      "celular": get("celular"), "email": get("email"),
                      "descricao": get("descricao"), "status": get("status") or "pendente",
                      "_linha": n})
    return itens, avisos


async def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    dry = "--dry-run" in sys.argv
    if not args:
        print('Uso: py importar_planilha.py "ARQUIVO.xlsx|.csv" [--dry-run]')
        sys.exit(1)
    path = Path(args[0])
    if not path.exists():
        print(f"Arquivo nao encontrado: {path}")
        sys.exit(1)
    ext = path.suffix.lower()
    matriz = ler_excel(path) if ext in (".xlsx", ".xls") else ler_csv(path)
    itens, avisos = linhas_para_itens(matriz)
    for a in avisos:
        print(f"[aviso] {a}")
    print(f"Linhas de dados: {len(itens)}")
    await init_db()
    destino = DATABASE_URL[:25] + "..." if USE_POSTGRES else DB_PATH
    print(f"Banco: {'Postgres' if USE_POSTGRES else 'SQLite'} ({destino}){' — DRY-RUN' if dry else ''}\n")
    ok, erros, vistos = 0, [], set()
    for it in itens:
        lin = it.pop("_linha")
        try:
            nome = it["destinatario"].strip()
            if not nome:
                raise ValueError(f"linha {lin}: destinatario vazio")
            codigo = validar_codigo(it["codigo"])
            email = validar_email(it["email"])
            celular = normalizar_celular(it["celular"])
            if not email and not celular:
                raise ValueError(f"linha {lin}: informe ao menos e-mail ou WhatsApp")
            status = (it["status"] or "pendente").strip().lower()
            if status not in ("pendente", "a_caminho", "entregue"):
                status = "pendente"
        except ValueError as e:
            erros.append(str(e))
            print(f"[ERRO] {e}")
            continue
        if codigo in vistos:
            erros.append(f"linha {lin}: duplicado na planilha ({codigo})")
            print(f"[ERRO] linha {lin}: duplicado na planilha ({codigo})")
            continue
        vistos.add(codigo)
        existe = await fetch_one("SELECT id FROM envios WHERE codigo = ?", (codigo,))
        if existe:
            erros.append(f"linha {lin}: {codigo} ja cadastrado (id={existe['id']})")
            print(f"[IGNORADO] linha {lin}: {codigo} ja cadastrado")
            continue
        if dry:
            print(f"[OK] linha {lin}: {nome} | {codigo} | {celular or '-'} | {email or '-'} | {status}")
            ok += 1
            continue
        _, last = await execute(
            "INSERT INTO envios (destinatario,codigo,celular,email,descricao,status,criado_por) VALUES (?,?,?,?,?,?,?)",
            (nome, codigo, celular, email, it["descricao"].strip(), status, None))
        print(f"[GRAVADO] id={last} | {nome} | {codigo} | {status}")
        ok += 1
    print(f"\nResultado: {ok} gravado(s), {len(erros)} erro(s)/ignorado(s) de {len(itens)} linha(s).")


if __name__ == "__main__":
    asyncio.run(main())

