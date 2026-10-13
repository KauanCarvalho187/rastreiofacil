# -*- coding: utf-8 -*-
"""Teste operacional de ponta a ponta (fluxo do dia a dia).

Sobe um servidor real (uvicorn) com banco TEMPORARIO e simula o uso normal:
  1. cria o admin (1o uso) e faz login (testa senha errada tambem);
  2. cadastra varias pessoas de teste (em lote);
  3. tenta duplicar um codigo (deve ser ignorado);
  4. muda o status de alguns envios;
  5. apaga alguns (lixeira), restaura um e purga outro;
  6. valida estatisticas e o total da listagem;
  7. salva a lista em Excel (envios_teste.xlsx), igual ao botao "Exportar Excel".

Uso:
    python teste_operacional.py            # sobe o servidor sozinho (porta livre)
    python teste_operacional.py --base http://127.0.0.1:3001   # usa um servidor ja rodando
"""
import argparse
import json
import os
import socket
import subprocess
import sys
import tempfile
import time
import urllib.error
import urllib.request

try:  # garante acentos no console do Windows
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

ROOT = os.path.dirname(os.path.abspath(__file__))
ADMIN = {"nome": "Admin Teste", "email": "admin@teste.com", "senha": "admin12345"}

PESSOAS = [
    ("Maria Silva",     "QX100000001BR", "(11) 98765-4321", "maria.silva@teste.com",  "Pedido #1001"),
    ("João Souza",      "QX100000002BR", "(11) 99123-4567", "joao.souza@teste.com",   "Pedido #1002"),
    ("Ana Paula",       "QX100000003BR", "(21) 98888-7777", "ana.paula@teste.com",    "Pedido #1003"),
    ("Carlos Eduardo",  "QX100000004BR", "(31) 97777-6666", "carlos.edu@teste.com",   "Presente"),
    ("Fernanda Lima",   "QX100000005BR", "(41) 96666-5555", "",                       "Sem e-mail (so WhatsApp)"),
    ("Rafael Costa",    "QX100000006BR", "(51) 95555-4444", "rafael.costa@teste.com", "Pedido #1006"),
    ("Beatriz Alves",   "QX100000007BR", "(61) 94444-3333", "bia.alves@teste.com",    "Pedido #1007"),
    ("Lucas Pereira",   "QX100000008BR", "(71) 93333-2222", "lucas.pereira@teste.com", "Pedido #1008"),
]

OK = 0
FALHA = 0


def check(nome, cond, detalhe=""):
    global OK, FALHA
    if cond:
        OK += 1
        print(f"  PASS - {nome}")
    else:
        FALHA += 1
        print(f"  FAIL - {nome}  {detalhe}")
    return cond


def _req(base, metodo, caminho, corpo=None, token=None):
    data = json.dumps(corpo).encode("utf-8") if corpo is not None else None
    req = urllib.request.Request(base + caminho, data=data, method=metodo)
    req.add_header("Content-Type", "application/json")
    if token:
        req.add_header("Authorization", "Bearer " + token)
    try:
        with urllib.request.urlopen(req, timeout=15) as r:
            b = r.read().decode("utf-8")
            return r.status, (json.loads(b) if b else {})
    except urllib.error.HTTPError as e:
        b = e.read().decode("utf-8")
        return e.code, (json.loads(b) if b else {})


def _porta_livre():
    with socket.socket() as s:
        s.bind(("127.0.0.1", 0))
        return s.getsockname()[1]


def _sobe_servidor():
    """Sobe o uvicorn com banco temporario; devolve (proc, base)."""
    pasta = tempfile.mkdtemp(prefix="rastreio-op-")
    porta = _porta_livre()
    env = os.environ.copy()
    env["DB_PATH"] = os.path.join(pasta, "teste.db")
    env["DATABASE_URL"] = ""
    env["RASTREIO_API_KEY"] = ""
    env["SMTP_USER"] = ""
    env["SMTP_PASS"] = ""
    proc = subprocess.Popen(
        [sys.executable, "-m", "uvicorn", "main:app", "--host", "127.0.0.1", "--port", str(porta)],
        cwd=ROOT, env=env, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    base = f"http://127.0.0.1:{porta}"
    for _ in range(60):
        try:
            if _req(base, "GET", "/api/health")[0] == 200:
                return proc, base
        except Exception:
            time.sleep(0.5)
    proc.terminate()
    raise RuntimeError("servidor nao subiu")


def _lista(base, token, status=""):
    q = f"?status={status}&limit=500" if status else "?limit=500"
    return _req(base, "GET", "/api/envios" + q, token=token)[1]


def _salvar_excel(caminho, envios):
    """Mesmas colunas do botao 'Exportar Excel' do sistema."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Envios"
    cols = ["Nome", "Codigo", "Celular", "Email", "Descricao", "Status"]
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2563EB")
    for e in envios:
        ws.append([e["destinatario"], e["codigo"], e["celular"], e["email"], e["descricao"], e["status"]])
    for col, largura in zip("ABCDEF", (22, 16, 18, 28, 24, 12)):
        ws.column_dimensions[col].width = largura
    wb.save(caminho)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--base", default="", help="usa um servidor ja rodando (ex.: http://127.0.0.1:3001)")
    ap.add_argument("--excel", default=os.path.join(ROOT, "envios_teste.xlsx"))
    args = ap.parse_args()

    proc = None
    if args.base:
        base = args.base.rstrip("/")
        print(f"Usando servidor existente em {base}\n")
    else:
        print("Subindo servidor de teste (banco temporario)...")
        proc, base = _sobe_servidor()
        print(f"Servidor no ar em {base}\n")

    try:
        # 1) servico + login -------------------------------------------------
        print("== 1. Servico e login ==")
        st, d = _req(base, "GET", "/api/health")
        check("health responde ok", st == 200 and d.get("ok") is True)

        st, d = _req(base, "POST", "/api/auth/registrar", ADMIN)
        check("cria admin no 1o uso (sem token)",
              st == 200 and d.get("admin") is True, f"status={st} resp={d}")

        st, d = _req(base, "POST", "/api/auth/login", {"email": ADMIN["email"], "senha": "errada"})
        check("login com senha errada -> 401", st == 401, f"status={st}")

        st, d = _req(base, "POST", "/api/auth/login", {"email": ADMIN["email"], "senha": ADMIN["senha"]})
        token = d.get("token", "")
        check("login com senha certa devolve token", st == 200 and bool(token))
        check("usuario logado e admin", d.get("user", {}).get("admin") is True)

        # 2) cadastro de varias pessoas -------------------------------------
        print("\n== 2. Cadastro de pessoas de teste ==")
        corpo = {"envios": [
            {"destinatario": n, "codigo": c, "celular": cel, "email": mail, "descricao": desc}
            for (n, c, cel, mail, desc) in PESSOAS]}
        st, d = _req(base, "POST", "/api/envios", corpo, token)
        check(f"cadastra as {len(PESSOAS)} pessoas em lote",
              st == 200 and len(d.get("envios", [])) == len(PESSOAS) and not d.get("erros"),
              f"status={st} criados={len(d.get('envios', []))} erros={d.get('erros')}")

        # 3) duplicado deve ser ignorado ------------------------------------
        print("\n== 3. Duplicidade ==")
        st, d = _req(base, "POST", "/api/envios",
                     {"destinatario": "Clone", "codigo": PESSOAS[0][1], "email": "x@teste.com"}, token)
        check("codigo repetido e ignorado (nao cria)",
              st == 200 and len(d.get("envios", [])) == 0 and any("cadastrado" in e for e in d.get("erros", [])),
              f"resp={d}")

        # 4) muda status -----------------------------------------------------
        print("\n== 4. Mudanca de status ==")
        todos = _lista(base, token)["envios"]
        por_codigo = {e["codigo"]: e["_id"] for e in todos}
        _req(base, "PATCH", f"/api/envios/{por_codigo['QX100000002BR']}/status", {"status": "a_caminho"}, token)
        _req(base, "PATCH", f"/api/envios/{por_codigo['QX100000003BR']}/status", {"status": "a_caminho"}, token)
        st, d = _req(base, "PATCH", f"/api/envios/{por_codigo['QX100000004BR']}/status", {"status": "entregue"}, token)
        check("status muda para entregue", st == 200 and d.get("status") == "entregue")
        st, _ = _req(base, "PATCH", f"/api/envios/{por_codigo['QX100000002BR']}/status", {"status": "voando"}, token)
        check("status invalido -> 422", st == 422, f"status={st}")

        # 5) apagar alguns (lixeira) + restaurar + purgar --------------------
        print("\n== 5. Apagar (lixeira), restaurar e purgar ==")
        apagar = ["QX100000005BR", "QX100000006BR", "QX100000007BR"]
        for cod in apagar:
            st, d = _req(base, "DELETE", f"/api/envios/{por_codigo[cod]}", token=token)
            check(f"apaga {cod} -> lixeira", st == 200 and "lixeira" in d.get("message", ""))

        ativos = _lista(base, token)
        lixo = _lista(base, token, status="excluido")
        check("ativos caem para 5", ativos["total"] == 5, f"total={ativos['total']}")
        check("lixeira tem 3", lixo["total"] == 3, f"total={lixo['total']}")

        st, d = _req(base, "POST", f"/api/envios/{por_codigo['QX100000005BR']}/restaurar", {}, token)
        check("restaura 1 da lixeira", st == 200, f"status={st} resp={d}")
        st, d = _req(base, "DELETE", f"/api/envios/{por_codigo['QX100000006BR']}/purga", token=token)
        check("purga 1 definitivamente", st == 200, f"status={st} resp={d}")

        ativos = _lista(base, token)
        lixo = _lista(base, token, status="excluido")
        check("ativos voltam para 6", ativos["total"] == 6, f"total={ativos['total']}")
        check("lixeira fica com 1", lixo["total"] == 1, f"total={lixo['total']}")

        # 6) estatisticas ----------------------------------------------------
        print("\n== 6. Estatisticas ==")
        st, e = _req(base, "GET", "/api/estatisticas", token=token)
        check("estatisticas somam certo",
              e.get("total") == 7 and e.get("entregue") == 1 and e.get("a_caminho") == 2 and e.get("pendente") == 3,
              f"resp={e}")

        # 7) exportar Excel --------------------------------------------------
        print("\n== 7. Exportar Excel ==")
        envios = _lista(base, token)["envios"]  # recorte ativo (sem lixeira), como no botao do sistema
        _salvar_excel(args.excel, envios)
        existe = os.path.exists(args.excel)
        check("arquivo Excel gerado", existe and os.path.getsize(args.excel) > 0, args.excel)

        print("\n== Pessoas ativas exportadas ==")
        for e in sorted(envios, key=lambda x: x["codigo"]):
            print(f"   {e['codigo']}  {e['destinatario']:<16} {e['status']:<10} {e['email'] or e['celular']}")
        print(f"\n   Excel: {args.excel}")

    finally:
        if proc:
            proc.terminate()
            try:
                proc.wait(timeout=10)
            except Exception:
                proc.kill()

    print(f"\nRESULTADO: {OK} passou, {FALHA} falhou")
    return 1 if FALHA else 0


if __name__ == "__main__":
    sys.exit(main())
